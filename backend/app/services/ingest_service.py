import uuid
import csv
from datetime import datetime, timezone
from pathlib import Path

import aiofiles
import httpx
from fastapi import HTTPException, UploadFile
from openpyxl import load_workbook
from pypdf import PdfReader
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.config import get_settings
from app.models.db_models import FileRecord
from app.services.dashscope_client import DashScopeClient
from app.services.vector_store import VectorStore


class IngestService:
    allowed_extensions = {".pdf", ".txt", ".md", ".csv", ".xlsx"}
    max_chunks = 30000

    def __init__(self) -> None:
        self.settings = get_settings()
        self.embed_batch_size = self.settings.embedding_batch_size
        self.max_chunks = self.settings.max_chunks_per_file
        self.vector_store = VectorStore()
        self.embedder = DashScopeClient()

    async def save_and_ingest(self, upload_file: UploadFile, db: AsyncSession) -> FileRecord:
        suffix = Path(upload_file.filename or "").suffix.lower()
        if suffix not in self.allowed_extensions:
            raise HTTPException(status_code=400, detail="Only PDF/TXT/MD/CSV/XLSX are supported.")

        file_id = str(uuid.uuid4())
        save_path = Path(self.settings.upload_dir) / f"{file_id}_{upload_file.filename}"
        file_size = 0
        max_bytes = self.settings.max_upload_mb * 1024 * 1024
        async with aiofiles.open(save_path, "wb") as f:
            while True:
                chunk = await upload_file.read(1024 * 1024)
                if not chunk:
                    break
                file_size += len(chunk)
                if file_size > max_bytes:
                    raise HTTPException(status_code=413, detail=f"File too large, max {self.settings.max_upload_mb}MB.")
                await f.write(chunk)

        chunks_iter = self._iter_chunks(save_path)
        first_chunk = next(chunks_iter, None)
        if not first_chunk:
            raise HTTPException(status_code=400, detail="Document has no valid text chunks.")

        chunk_count = 0
        batch: list[dict] = [first_chunk]
        for chunk in chunks_iter:
            batch.append(chunk)
            if len(batch) >= self.embed_batch_size:
                chunk_count += await self._embed_and_upsert_batch(
                    batch=batch,
                    file_id=file_id,
                    file_name=upload_file.filename or "",
                    start_index=chunk_count,
                )
                if chunk_count > self.max_chunks:
                    raise HTTPException(status_code=400, detail="Document too large after chunking.")
                batch = []
        if batch:
            chunk_count += await self._embed_and_upsert_batch(
                batch=batch,
                file_id=file_id,
                file_name=upload_file.filename or "",
                start_index=chunk_count,
            )

        record = FileRecord(
            id=uuid.UUID(file_id),
            file_name=upload_file.filename or "",
            file_size=file_size,
            status="processed",
            chunk_count=chunk_count,
        )
        db.add(record)
        await db.commit()
        await db.refresh(record)
        return record

    def _parse_text(self, path: Path) -> str:
        suffix = path.suffix.lower()
        if suffix in {".txt", ".md"}:
            return path.read_text(encoding="utf-8", errors="ignore")
        if suffix == ".csv":
            return self._parse_csv_text(path)
        if suffix == ".xlsx":
            return self._parse_xlsx_text(path)
        if suffix == ".pdf":
            reader = PdfReader(str(path))
            return "\n".join(page.extract_text() or "" for page in reader.pages)
        return ""

    def _iter_chunks(self, path: Path):
        suffix = path.suffix.lower()
        if suffix in {".txt", ".md"}:
            text = path.read_text(encoding="utf-8", errors="ignore")
            yield from self._iter_chunks_from_text(text)
            return
        if suffix == ".csv":
            text = self._parse_csv_text(path)
            yield from self._iter_chunks_from_text(text)
            return
        if suffix == ".xlsx":
            text = self._parse_xlsx_text(path)
            yield from self._iter_chunks_from_text(text)
            return
        if suffix == ".pdf":
            reader = PdfReader(str(path))
            running_offset = 0
            total_chars = 0
            if len(reader.pages) > self.settings.max_pdf_pages:
                raise HTTPException(status_code=400, detail=f"PDF pages exceed limit: {self.settings.max_pdf_pages}")
            for page in reader.pages:
                text = page.extract_text() or ""
                total_chars += len(text)
                if total_chars > self.settings.max_pdf_chars:
                    raise HTTPException(
                        status_code=400,
                        detail=f"PDF text too large after extraction: {self.settings.max_pdf_chars} chars limit.",
                    )
                for chunk in self._iter_chunks_from_text(text):
                    yield {
                        "text": chunk["text"],
                        "start_offset": running_offset + chunk["start_offset"],
                        "end_offset": running_offset + chunk["end_offset"],
                    }
                running_offset += len(text)

    async def _embed_and_upsert_batch(self, batch: list[dict], file_id: str, file_name: str, start_index: int) -> int:
        chunk_texts = [item["text"][: self.settings.embedding_max_chunk_chars] for item in batch]
        vectors = await self._safe_embed_texts(chunk_texts)
        ids = [f"{file_id}_{start_index + i}" for i in range(len(batch))]
        metadata = [
            {
                "file_id": file_id,
                "file_name": file_name,
                "chunk_index": start_index + i,
                "start_offset": batch[i]["start_offset"],
                "end_offset": batch[i]["end_offset"],
                "upload_time": datetime.now(timezone.utc).isoformat(),
            }
            for i in range(len(batch))
        ]
        self.vector_store.upsert(ids=ids, embeddings=vectors, documents=chunk_texts, metadatas=metadata)
        return len(batch)

    async def _safe_embed_texts(self, chunk_texts: list[str]) -> list[list[float]]:
        if not chunk_texts:
            return []
        try:
            vectors, _ = await self.embedder.embed_texts(chunk_texts)
            return vectors
        except (httpx.ReadError, httpx.TimeoutException, httpx.ConnectError):
            if len(chunk_texts) == 1:
                raise HTTPException(status_code=502, detail="Embedding API unstable, please retry later.")
            mid = len(chunk_texts) // 2
            left = await self._safe_embed_texts(chunk_texts[:mid])
            right = await self._safe_embed_texts(chunk_texts[mid:])
            return left + right
        except httpx.HTTPStatusError as exc:
            if exc.response.status_code == 400:
                if len(chunk_texts) == 1:
                    raise HTTPException(
                        status_code=400,
                        detail="Embedding API rejected this chunk content. Please split document or reduce chunk size.",
                    ) from exc
                mid = len(chunk_texts) // 2
                left = await self._safe_embed_texts(chunk_texts[:mid])
                right = await self._safe_embed_texts(chunk_texts[mid:])
                return left + right
            raise HTTPException(status_code=502, detail=f"Embedding API failed: {exc.response.status_code}") from exc

    def _iter_chunks_from_text(self, text: str):
        if not text or not text.strip():
            return

        size = self.settings.chunk_size
        overlap = self.settings.chunk_overlap
        cursor = 0
        text_len = len(text)
        while cursor < text_len:
            end = min(cursor + size, text_len)
            window = text[cursor:end]
            split_idx = max(window.rfind("。"), window.rfind("."), window.rfind("\n"))
            if 80 < split_idx < len(window) - 1:
                window = window[: split_idx + 1]
                end = cursor + split_idx + 1
            chunk = window.strip()
            if chunk:
                yield {
                    "text": chunk,
                    "start_offset": cursor,
                    "end_offset": end,
                }
            if end >= text_len:
                break
            # Ensure forward progress even when overlap is larger than the effective window.
            next_cursor = max(end - overlap, 0)
            cursor = next_cursor if next_cursor > cursor else cursor + 1

    def _parse_csv_text(self, path: Path) -> str:
        rows: list[str] = []
        with path.open("r", encoding="utf-8", errors="ignore", newline="") as fp:
            reader = csv.reader(fp)
            headers: list[str] | None = None
            for idx, row in enumerate(reader):
                if idx == 0:
                    headers = [str(cell).strip() for cell in row]
                    continue
                if not any(str(cell).strip() for cell in row):
                    continue
                if headers and len(headers) == len(row):
                    parts = [f"{headers[i]}={str(row[i]).strip()}" for i in range(len(row))]
                    rows.append(" | ".join(parts))
                else:
                    rows.append(" | ".join(str(cell).strip() for cell in row if str(cell).strip()))
        return "\n".join(rows)

    def _parse_xlsx_text(self, path: Path) -> str:
        wb = load_workbook(filename=str(path), read_only=True, data_only=True)
        lines: list[str] = []
        for sheet in wb.worksheets:
            rows = sheet.iter_rows(values_only=True)
            headers_row = next(rows, None)
            headers = [str(cell).strip() if cell is not None else f"col_{idx+1}" for idx, cell in enumerate(headers_row or [])]
            for row in rows:
                values = ["" if cell is None else str(cell).strip() for cell in row]
                if not any(values):
                    continue
                pairs = []
                for idx, value in enumerate(values):
                    if not value:
                        continue
                    key = headers[idx] if idx < len(headers) and headers[idx] else f"col_{idx+1}"
                    pairs.append(f"{key}={value}")
                if pairs:
                    lines.append(f"[sheet={sheet.title}] " + " | ".join(pairs))
        wb.close()
        return "\n".join(lines)
